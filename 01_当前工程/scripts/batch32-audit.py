from pathlib import Path
import re,json,hashlib,zipfile
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
r=Path(__file__).resolve().parents[1];b=r.parent/'Batch3.2_备份_v0.3.1_20260906';out=r/'artifacts/batch3-2'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
old=(b/'docs/PRD_v0.6.md').read_text(encoding='utf-8');new=(r/'docs/PRD_v0.6.md').read_text(encoding='utf-8')
parts=lambda s:re.split(r'(?=^## \d{2}｜)',s,flags=re.M)[1:]
assert len(parts(new))==19
unchanged=[]
for i,(a,c) in enumerate(zip(parts(old),parts(new)),1):
 if i not in [1,5,7,12]:assert a==c,i;unchanged.append(i)
codes=re.findall(r'```text[\s\S]*?```',old)
assert len(codes)==3 and all(code in new for code in codes)
frozen=[]
for folder in ['src','specs','data','design','.agents']:
 for p in (b/folder).rglob('*'):
  if p.is_file() and str(p.relative_to(b)).replace('\\','/') not in ['src/Reader.tsx','src/main.tsx']:
   assert sha(p)==sha(r/p.relative_to(b)),str(p);frozen.append(p.relative_to(b).as_posix())
for f in ['pnpm-lock.yaml','docs/PRD_v0.5.1.md','docs/AI_CONTRACT.md','docs/IMPLEMENTATION_CONTRACT.md']:
 assert sha(b/f)==sha(r/f),f
old_images=list(b.rglob('*.png'))
for p in old_images:assert sha(p)==sha(r/p.relative_to(b)),str(p)
for f in ['docs/HANDOFF.md','docs/APPROVAL_STATUS.md','README.md']:
 assert (b/f).read_text(encoding='utf-8') in (r/f).read_text(encoding='utf-8'),f
def plain(text):
 text=re.sub(r'^!\[([^\]]+)\]\([^)]*\)$',r'\1',text,flags=re.M)
 lines=[line for line in text.splitlines() if not re.match(r'^```|^\|[\s|:-]+\|$|^!\[',line)]
 text='\n'.join(re.sub(r'^#{1,6} |^> ?','',line) for line in lines)
 text=re.sub(r'\[([^\]]+)\]\([^)]*\)',r'\1',text)
 return re.sub(r'\s|`|\||\*\*','',text)
docpath=r/'public/documents/PRD_v0.6_需求评审稿.docx';doc=Document(docpath)
word=re.sub(r'\s','', ''.join(t.text or '' for t in doc.element.body.iter(qn('w:t'))))
expected=plain(new)
if word!=expected:
 index=next((i for i,(a,c) in enumerate(zip(word,expected)) if a!=c),min(len(word),len(expected)))
 raise AssertionError({'wordChars':len(word),'mdChars':len(expected),'at':index,'word':word[max(0,index-80):index+140],'md':expected[max(0,index-80):index+140]})
assert len(doc.tables)==22 and len(doc.inline_shapes)==1
for s in doc.sections:
 for p in s.header.paragraphs+s.footer.paragraphs:assert '提交候选' not in p.text
pdf_results={}
for name,path in [('browser',r/'public/documents/PRD_v0.6_需求评审稿.pdf'),('word',out/'word-render/word.pdf')]:
 pdf=PdfReader(path);text=re.sub(r'\s','', ''.join(p.extract_text() for p in pdf.pages))
 for part in parts(new):assert re.sub(r'\s','',part.splitlines()[0].removeprefix('## ')) in text,(name,part[:40])
 for s in ['Q01','Q02','Q03','Q04','Q05','positiveMatch','negativeMatch','0.30','0.20','0.60','不代表网易云正式推荐算法','状态：需求评审稿']:assert s in text,(name,s)
 for s in ['提交候选','负责人待复核','NOT_RUN','ERR_BLOCKED','Batch3']:assert s not in text,(name,s)
 pdf_results[name]={'pages':len(pdf.pages),'all19TitlesAndParametersAndQuestions':True,'noSupersededStatus':True}
report={'status':'PASS','application':'0.3.2','prd':'0.6','state':'需求评审稿','unchangedSections':unchanged,'unchangedCodeBlocks':3,'tables':22,'figure':1,'wordEntireBodyMatchesMd':True,'mdSha256':sha(r/'docs/PRD_v0.6.md'),'wordSha256':sha(docpath),'oldPngPreserved':len(old_images),'frozenFiles':frozen,'dependenciesUnchanged':True,'oldHistoryPreserved':True,'pdf':pdf_results}
(out/'final-audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='frozenFiles'},ensure_ascii=False,indent=2))
